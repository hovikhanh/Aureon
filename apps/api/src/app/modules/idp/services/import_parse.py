"""C5 workbook parsing (EP-DC-17/19, R-DEC-62) — streaming, limit-checked, no DB.

`.xlsx` is opened with openpyxl `read_only=True, data_only=True` straight from the stored file
(rows are streamed, the workbook is never materialised); `.csv` must be UTF-8 (strict decode —
a non-UTF-8 file is `422`, never silently "replaced"). Limits (`IDP_IMPORT_MAX_*`) → `413
FILE_TOO_LARGE` with `details.limit = rows | sheets | columns`.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from app.core.config import get_config
from app.core.errors import AppError
from app.core.errors import ValidationError as AppValidationError

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV = "text/csv"
HEADER_SCAN_ROWS = 50  # header_row ∈ 1..50 (05 §2.7 CHECK)
HEADER_ALLOWANCE = 5  # non-data rows (title, blank, header) tolerated above the row limit


@dataclass
class Sheet:
    name: str
    index: int  # 1-based sheet position (= DocumentPage.page_number)
    rows: list[tuple[int, list[Any]]] = field(default_factory=list)  # (Excel row no, cells)

    def header_row(self) -> int:
        """Default header = the widest row (most non-empty cells, first wins) within the first
        50 — skips title/blank rows above a table (DOC-06 "Danh sách nhân sự …" banner)."""
        best, width = 1, 0
        for number, cells in self.rows:
            if number > HEADER_SCAN_ROWS:
                break
            filled = sum(1 for c in cells if not _blank(c))
            if filled > width:
                best, width = number, filled
        return best

    def headers(self, header_row: int) -> list[str]:
        row = next((c for n, c in self.rows if n == header_row), [])
        return [cell_text(c) for c in row]

    def data(self, header_row: int) -> list[tuple[int, list[Any]]]:
        return [(n, c) for n, c in self.rows if n > header_row and any(not _blank(x) for x in c)]


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return (
            value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
        )
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def json_cell(value: Any) -> Any:
    """JSON-safe raw cell (dates → ISO text, numbers kept)."""
    if isinstance(value, (datetime, date)):
        return cell_text(value)
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    return str(value)


def too_large(limit: str, maximum: int) -> AppError:
    return AppError(f"import exceeds the {limit} limit ({maximum})", code="FILE_TOO_LARGE",
                    status_code=413, details={"limit": limit, "max": maximum})  # fmt: skip


def parse(path: Path, mime: str) -> list[Sheet]:
    cfg = get_config()
    limits = (cfg.idp_import_max_rows, cfg.idp_import_max_sheets, cfg.idp_import_max_columns)
    if mime == CSV:
        return [_csv(path, *limits)]
    return _xlsx(path, *limits)


def _add_row(sheet: Sheet, number: int, cells: list[Any], max_rows: int, max_cols: int) -> None:
    while cells and _blank(cells[-1]):
        cells.pop()
    if not cells:
        return
    if len(cells) > max_cols:
        raise too_large("columns", max_cols)
    sheet.rows.append((number, cells))
    if len(sheet.rows) > max_rows + HEADER_ALLOWANCE:  # header + title rows above the table
        raise too_large("rows", max_rows)


def _xlsx(path: Path, max_rows: int, max_sheets: int, max_cols: int) -> list[Sheet]:
    from openpyxl import load_workbook

    fh = path.open("rb")  # stored key has no extension → openpyxl needs a file object
    try:
        wb = load_workbook(fh, read_only=True, data_only=True)
    except Exception as exc:  # corrupt/encrypted OOXML → user error, not 500
        fh.close()
        raise AppValidationError("unreadable workbook", details={"reason": "unreadable"}) from exc
    try:
        if len(wb.worksheets) > max_sheets:
            raise too_large("sheets", max_sheets)
        sheets: list[Sheet] = []
        for idx, ws in enumerate(wb.worksheets, start=1):
            sheet = Sheet(ws.title, idx)
            for number, row in enumerate(ws.iter_rows(values_only=True), start=1):
                _add_row(sheet, number, list(row), max_rows, max_cols)
            sheets.append(sheet)
        return sheets
    finally:
        wb.close()
        fh.close()


def _csv(path: Path, max_rows: int, _sheets: int, max_cols: int) -> Sheet:
    sheet = Sheet("CSV", 1)
    try:
        with path.open(encoding="utf-8-sig", errors="strict", newline="") as fh:
            for number, row in enumerate(csv.reader(fh), start=1):
                _add_row(sheet, number, list(row), max_rows, max_cols)
    except UnicodeDecodeError as exc:
        raise AppValidationError("CSV must be UTF-8", details={"reason": "encoding"}) from exc
    return sheet


def write_error_file(headers: list[str], rows: list[list[Any]]) -> bytes:
    """`.xlsx` (openpyxl `write_only`) — original columns + `Lỗi` (EP-DC-21)."""
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Lỗi")
    ws.append([*headers, "Lỗi"])
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
