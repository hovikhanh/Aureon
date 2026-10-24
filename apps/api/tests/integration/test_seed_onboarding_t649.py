"""T-649 / TC-M16-22 — UC-07 seed + sample file (08 §4.6, R-DEC-79).

* `data/ABC_DanhSachNhanSu_Onboarding.xlsx` (scripts/gen_demo_docs.py): the C5 preview splits
  exactly the valid / invalid rows the generator computes (unknown manager, bad date, duplicate
  email); ADM never gets the salary column (level 4).
* Seed: 6 agents, no autonomy policy row (registry defaults), `onboarding.state = completed` for
  ABC — idempotent, a `reset` row is left untouched by a re-seed.
"""

import uuid
from pathlib import Path
from types import ModuleType
from typing import Any

import pytest
from sqlalchemy import func, select

from app.modules.ai.models import AgentProfile, AutonomyPolicy
from app.modules.platform.services.onboarding_service import KEY, STEPS
from app.modules.platform.services.settings_service import SystemSettingService
from app.seed.ai_seed import seed_agents
from app.seed.onboarding_seed import seed_onboarding
from tests.conftest import ApiClient
from tests.integration.idp.conftest import idp_store  # noqa: F401 — re-exported fixture
from tests.integration.test_seed_v3 import AS_OF, TRUNCATE_ALL, _seed_all

pytestmark = pytest.mark.integration
ROOT = Path(__file__).resolve().parents[4]
SAMPLE = ROOT / "data" / "ABC_DanhSachNhanSu_Onboarding.xlsx"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _generator() -> ModuleType:
    """`uc07_expectations()` without reportlab: load only the constants/functions we need."""
    src = (ROOT / "scripts" / "gen_demo_docs.py").read_text(encoding="utf-8")
    start = src.index("# UC-07 — onboarding roster")
    end = src.index("def gen_uc07")
    mod = ModuleType("uc07")
    exec("from typing import Final\n" + src[start:end], mod.__dict__)  # noqa: S102 — repo file
    return mod


def test_sample_file_is_the_generator_output() -> None:
    from openpyxl import load_workbook

    exp = _generator().uc07_expectations()
    ws = load_workbook(SAMPLE, read_only=True).worksheets[0]
    head, *rows = list(ws.iter_rows(values_only=True))
    assert "Email" in head and "Lương" in head
    assert len(rows) == exp["rows_total"] and exp["rows_invalid"] == 3
    assert len({r[0] for r in rows[: exp["rows_valid"]]}) == exp["departments"]
    web = ROOT / "apps" / "web" / "public" / "templates" / SAMPLE.name
    assert web.read_bytes() == SAMPLE.read_bytes()  # the SCR-66 template link serves the same file


@pytest.mark.usefixtures("idp_store")
async def test_preview_splits_intentional_error_rows(api: ApiClient, db_session) -> None:
    exp = _generator().uc07_expectations()
    adm = uuid.uuid4()
    h = api.headers("ADM", idem=True, user_id=adm)
    r = await api.http.post("/api/v1/data-capture/imports", headers=h,
                            data={"document_type": "HR_ROSTER"},
                            files=[("file", (SAMPLE.name, SAMPLE.read_bytes(), XLSX))])  # fmt: skip
    assert r.status_code == 201, r.text
    p: dict[str, Any] = r.json()["data"]
    assert p["suggested_document_type"]["code"] == "HR_ROSTER"
    assert p["summary"] == {k: exp[k] for k in ("rows_total", "rows_valid", "rows_invalid")}
    mapped = {v["field"] for v in p["column_mapping"].values() if v["field"]}
    assert "email" in mapped and "salary" not in mapped  # ADM: level 4 column dropped
    bad = await api.http.get(f"/api/v1/data-capture/imports/{p['batch_id']}/rows",
                             headers=api.headers("ADM", user_id=adm),
                             params={"status": "invalid", "order": "asc"})  # fmt: skip
    assert bad.status_code == 200, bad.text
    codes = [{(e["field"], e["rule_code"]) for e in x["errors"]} for x in bad.json()["data"]]
    assert codes == [{("manager_code", "MASTER_NOT_FOUND")}, {("hire_date", "DATE_FORMAT")},
                     {("email", "DUPLICATE")}]  # fmt: skip
    assert all(
        not e["params"] for x in bad.json()["data"] for e in x["errors"] if e["field"] == "email"
    )  # level 3 email never echoed


@pytest.fixture
async def seeded(db_session: Any, monkeypatch: pytest.MonkeyPatch) -> Any:
    from app.core.config import get_config

    monkeypatch.setattr(get_config(), "demo_reference_date", AS_OF.isoformat())
    await db_session.execute(TRUNCATE_ALL)
    await _seed_all(db_session)
    await seed_agents(db_session)
    return db_session


async def test_seed_agents_no_policies_onboarding_completed(seeded: Any) -> None:
    assert await seed_onboarding(seeded, None) == "created"
    assert await seed_onboarding(seeded, None) == "unchanged"
    n_agents = (await seeded.execute(select(func.count()).select_from(AgentProfile))).scalar()
    n_pol = (await seeded.execute(select(func.count()).select_from(AutonomyPolicy))).scalar()
    assert (n_agents, n_pol) == (6, 0)
    row = await SystemSettingService(seeded).get(KEY)
    assert row is not None and row.version == 1
    v = row.value
    assert v["status"] == "completed" and v["started_at"] < v["completed_at"]
    assert {s: v["steps"][s]["status"] for s in STEPS} == dict.fromkeys(STEPS, "done")
    assert set(v["refs"]) == {"organization_id"}  # ids only (05 §2.8)


async def test_reseed_keeps_a_reset_state(seeded: Any) -> None:
    from app.core.deps import Principal
    from app.modules.auth.permission_matrix import ROLE_PERMISSIONS
    from app.modules.platform.schemas import OnboardingPatchIn
    from app.modules.platform.services.onboarding_service import OnboardingService

    await seed_onboarding(seeded, None)
    adm = Principal(user_id=uuid.uuid4(), role="ADM",
                    permissions=frozenset(ROLE_PERMISSIONS["ADM"]))  # fmt: skip
    svc = OnboardingService(seeded)
    s = await svc.get()
    assert s.status == "completed" and s.steps["review"].status == "done"
    out = await svc.apply(adm, OnboardingPatchIn(action="reset", reason="e2e", version=s.version))
    assert out.status == "not_started"
    assert await seed_onboarding(seeded, None) == "unchanged"
    assert (await svc.get()).status == "not_started"
