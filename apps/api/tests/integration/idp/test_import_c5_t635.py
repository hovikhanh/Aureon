"""T-635 / TC-M15-05/06 — C5 import EP-DC-17..22 (R-DEC-62/65/66/76).

The workbook mirrors DOC-06 (sheet "Nhân sự", 8 mapped columns + salary) with intentional error
rows: missing department, bad date, unknown manager. Built in memory with openpyxl.
"""

import io
import tracemalloc
import uuid
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, update

from app.modules.idp.models import ExtractedField, ImportBatch, ReviewAssignment, SourceDocument
from tests.conftest import ApiClient

pytestmark = pytest.mark.integration

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEAD = ["Mã phòng ban", "Phòng ban", "Chức danh", "Mã NV", "Họ tên", "Quản lý trực tiếp",
        "Ngày vào làm", "Trạng thái", "Lương"]  # fmt: skip
GOOD = [
    ["PB-T1", "Phòng Thử", "Trưởng phòng", "T635-01", "Nguyễn A", None, "01/02/2024", "active",
     25000000],
    ["PB-T1", "Phòng Thử", "Chuyên viên", "T635-02", "Trần B", "T635-01", "15/03/2025", "active",
     15000000],
]  # fmt: skip
BAD = [
    [None, None, "Chuyên viên", "T635-03", "Lê C", "T635-01", "01/01/2025", "active", 1],
    ["PB-T1", "Phòng Thử", "Chuyên viên", "T635-04", "Phạm D", "T635-01", "31/13/2025",
     "active", 1],
    ["PB-T1", "Phòng Thử", "Chuyên viên", "T635-05", "Võ E", "NV-KHONG-CO", "01/01/2025",
     "active", 1],
]  # fmt: skip


def _xlsx(
    rows: list[list[Any]], *, sheets: int = 1, head: list[str] = HEAD, tag: str = ""
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Nhân sự"
    ws.append([f"Danh sách nhân sự {tag}"])  # title row → header detected on row 3
    ws.append([])
    ws.append(head)
    for r in rows:
        ws.append(r)
    for i in range(1, sheets):
        wb.create_sheet(f"S{i}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _upload(api: ApiClient, role: str, content: bytes, *, name: str = "ds.xlsx",
                  mime: str = XLSX, **kw: Any) -> Any:  # fmt: skip
    h = api.headers(role, idem=True, **kw)
    return await api.http.post("/api/v1/data-capture/imports", headers=h,
                               files=[("file", (name, content, mime))])  # fmt: skip


async def test_preview_maps_columns_and_flags_bad_rows(
    api: ApiClient, db_session, idp_store
) -> None:
    r = await _upload(api, "CEO", _xlsx(GOOD + BAD, tag=uuid.uuid4().hex))
    assert r.status_code == 201, r.text
    p = r.json()["data"]
    assert p["suggested_document_type"]["code"] == "HR_ROSTER"
    assert p["sheets"][0]["name"] == "Nhân sự" and p["sheets"][0]["header_row"] == 3
    mapped = {k: v["field"] for k, v in p["column_mapping"].items() if v["field"]}
    assert len(mapped) == 9 and mapped["Lương"] == "salary"  # CEO sees level 4
    assert p["summary"] == {"rows_total": 5, "rows_valid": 2, "rows_invalid": 3}
    bid = p["batch_id"]
    bad = await api.call("GET", f"/data-capture/imports/{bid}/rows", "CEO",
                         params={"status": "invalid", "order": "asc"})  # fmt: skip
    rows = bad.json()["data"]
    assert [x["row_number"] for x in rows] == [6, 7, 8]
    codes = [{e["rule_code"] for e in x["errors"]} for x in rows]
    assert "REQUIRED" in codes[0] and codes[1] == {"DATE_FORMAT"}
    assert codes[2] == {"MASTER_NOT_FOUND"}
    assert bad.json()["meta"]["pagination"]["total"] == 3
    again = await api.call("GET", f"/data-capture/imports/{bid}", "CEO")
    assert again.status_code == 200 and again.json()["data"]["version"] == p["version"]


async def test_mapping_update_and_level4_drop(api: ApiClient, db_session, idp_store) -> None:
    uid = uuid.uuid4()
    h = api.headers("ADMIN", idem=True, user_id=uid)
    r = await api.http.post(
        "/api/v1/data-capture/imports",
        headers=h,
        data={"document_type": "HR_ROSTER"},
        files=[("file", ("ds.xlsx", _xlsx(GOOD + BAD, tag="adm"), XLSX))],
    )
    assert r.status_code == 201, r.text
    p = r.json()["data"]
    assert p["column_mapping"]["Lương"]["field"] is None  # ADM: level-4 column dropped
    batch = (await db_session.execute(select(ImportBatch).where(
        ImportBatch.id == uuid.UUID(p["batch_id"])))).scalar_one()  # fmt: skip
    assert batch.column_mapping["warnings"][0] == {"code": "idp.warn.level4_dropped",
                                                   "columns": ["Lương"]}  # fmt: skip
    rows = await api.call("GET", f"/data-capture/imports/{p['batch_id']}/rows", "CEO")
    assert all("salary" not in x["values"] for x in rows.json()["data"])  # never stored
    url = f"/api/v1/data-capture/imports/{p['batch_id']}/mapping"
    body = {"document_type": "HR_ROSTER", "sheet": "Nhân sự", "header_row": 3,
            "mapping": {"Mã NV": None}, "version": p["version"]}  # fmt: skip
    h = api.headers("ADMIN", user_id=uid)
    miss = await api.http.patch(url, headers=h, json=body)
    assert miss.status_code == 422
    assert miss.json()["error"]["details"]["unmapped_required"] == ["employee_code"]
    other = await api.http.patch(url, headers=api.headers("ADMIN"), json={**body, "mapping": {}})
    assert other.status_code == 403  # submitter only
    ok = await api.http.patch(
        url, headers=h, json={**body, "mapping": {"Lương": "salary", "Chức danh": None}}
    )
    assert ok.status_code == 200, ok.text
    q = ok.json()["data"]
    assert q["version"] == p["version"] + 1 and q["column_mapping"]["Lương"]["field"] is None
    assert q["column_mapping"]["Chức danh"]["field"] is None
    stale = await api.http.patch(url, headers=h, json={**body, "mapping": {}})
    assert stale.status_code == 409


async def test_error_file_and_masking(api: ApiClient, db_session, idp_store) -> None:
    r = await _upload(api, "CEO", _xlsx(GOOD + BAD, tag=uuid.uuid4().hex))
    bid = r.json()["data"]["batch_id"]
    ceo = await api.call("GET", f"/data-capture/imports/{bid}/error-file", "CEO")
    assert ceo.status_code == 200 and ceo.headers["content-type"].startswith(XLSX)
    ws = load_workbook(io.BytesIO(ceo.content), read_only=True).active
    got = list(ws.iter_rows(values_only=True))
    assert got[0][-1] == "Lỗi" and "Lương" in got[0] and len(got) == 4
    dir_f = await api.call("GET", f"/data-capture/imports/{bid}/error-file", "CEO")
    assert dir_f.status_code == 200
    rev_rows = await api.call("GET", f"/data-capture/imports/{bid}/rows", "CEO",
                              params={"order": "asc"})  # fmt: skip
    assert rev_rows.json()["data"][0]["values"]["salary"] == 25000000
    good = await _upload(api, "CEO", _xlsx(GOOD, tag=uuid.uuid4().hex))
    none = await api.call("GET", f"/data-capture/imports/{good.json()['data']['batch_id']}"
                          "/error-file", "CEO")  # fmt: skip
    assert none.status_code == 404


async def test_upload_rejections(api: ApiClient, db_session, idp_store, monkeypatch) -> None:
    from app.core.config import get_config

    xls = await _upload(api, "CEO", b"\xd0\xcf\x11\xe0" + b"0" * 64, name="a.xls",
                        mime="application/vnd.ms-excel")  # fmt: skip
    assert xls.status_code == 415
    fake = await _upload(api, "CEO", b"not a zip", name="a.xlsx")
    assert fake.status_code == 415
    latin = await _upload(api, "CEO", "Mã NV;Họ tên\n".encode("utf-16"), name="a.csv",
                          mime="text/csv")  # fmt: skip
    assert latin.status_code == 422
    monkeypatch.setattr(get_config(), "idp_import_max_rows", 1)
    many = await _upload(api, "CEO", _xlsx(GOOD + BAD, tag="rows"))
    assert many.status_code == 413 and many.json()["error"]["details"]["limit"] == "rows"
    monkeypatch.setattr(get_config(), "idp_import_max_sheets", 2)
    sh = await _upload(api, "CEO", _xlsx(GOOD, sheets=3, tag="sheets"))
    assert sh.status_code == 413 and sh.json()["error"]["details"]["limit"] == "sheets"
    monkeypatch.setattr(get_config(), "idp_import_max_columns", 5)
    co = await _upload(api, "CEO", _xlsx(GOOD, tag="cols"))
    assert co.status_code == 413 and co.json()["error"]["details"]["limit"] == "columns"


async def test_submit_valid_rows_then_dir_o_commits(api: ApiClient, db_session, idp_store) -> None:
    from app.modules.hr.services.hr_service import HrService

    uid = uuid.uuid4()
    r = await _upload(api, "CEO", _xlsx(GOOD + BAD, tag=uuid.uuid4().hex), user_id=uid)
    p = r.json()["data"]
    url = f"/api/v1/data-capture/imports/{p['batch_id']}/submit"
    other = await api.http.post(url, headers=api.headers("CEO", idem=True),
                                json={"version": p["version"]})  # fmt: skip
    assert other.status_code == 403
    sub = await api.http.post(url, headers=api.headers("CEO", idem=True, user_id=uid),
                              json={"version": p["version"]})  # fmt: skip
    assert sub.status_code == 200, sub.text
    out = sub.json()["data"]
    assert (out["rows_submitted"], out["rows_excluded"]) == (2, 3)
    assert out["document"]["status"] == "NEEDS_REVIEW"
    doc_id = uuid.UUID(out["document"]["id"])
    fields = (await db_session.execute(select(ExtractedField).where(
        ExtractedField.field_name == "employee_code",
        ExtractedField.source_location["sheet"].astext == "Nhân sự",
    ).order_by(ExtractedField.row_index))).scalars().all()  # fmt: skip
    assert [f.row_index for f in fields][-2:] == [4, 5] and fields[-1].document_page_id
    ra = (await db_session.execute(select(ReviewAssignment).where(
        ReviewAssignment.source_document_id == doc_id))).scalar_one()  # fmt: skip
    assert ra.is_current
    again = await api.http.post(url, headers=api.headers("CEO", idem=True, user_id=uid),
                                json={"version": p["version"] + 1})  # fmt: skip
    assert again.status_code == 409  # no longer VALIDATING
    dir_o = await api.call("POST", f"/idp/documents/{doc_id}/review/approve", "HR_MANAGER", json={})
    assert dir_o.status_code == 200, dir_o.text
    assert dir_o.json()["data"]["document"]["status"] == "COMMITTED"
    hr = HrService(db_session)
    assert await hr.department_id_by_code("PB-T1") is not None
    assert await hr.employee_id_by_code("T635-02") is not None


async def test_zero_valid_rows_and_duplicates(api: ApiClient, db_session, idp_store) -> None:
    uid = uuid.uuid4()
    content = _xlsx(BAD, tag=uuid.uuid4().hex)
    p = (await _upload(api, "CEO", content, user_id=uid)).json()["data"]
    sub = await api.http.post(f"/api/v1/data-capture/imports/{p['batch_id']}/submit",
                              headers=api.headers("CEO", idem=True, user_id=uid),
                              json={"version": p["version"]})  # fmt: skip
    assert sub.status_code == 422
    assert sub.json()["error"]["details"]["reason"] == "no_valid_rows"
    dup = await _upload(api, "CEO", content)
    assert dup.status_code == 409 and dup.json()["error"]["code"] == "DUPLICATE_RESOURCE"
    assert dup.json()["error"]["details"]["duplicate_of_id"] == p["document_id"]
    await db_session.execute(update(SourceDocument).where(
        SourceDocument.id == uuid.UUID(p["document_id"])).values(status="ROLLED_BACK"))  # fmt: skip
    for d in (await db_session.execute(select(SourceDocument).where(
            SourceDocument.status == "DUPLICATE"))).scalars():  # fmt: skip
        d.status = "ROLLED_BACK"  # the rejected duplicate is not an original either
    await db_session.flush()
    fresh = await _upload(api, "CEO", content)
    assert fresh.status_code == 201, fresh.text  # R-DEC-66: rolled back ≠ duplicate


async def test_5000_rows_stream_under_64mb(api: ApiClient, db_session, idp_store) -> None:
    rows = [["PB-M", "Phòng M", "CV", f"M{i:05d}", f"Người {i}", None, "01/02/2024", "active",
             1000000] for i in range(5000)]  # fmt: skip
    content = _xlsx(rows, tag="mem")
    tracemalloc.start()
    r = await _upload(api, "CEO", content)
    _, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    assert r.status_code == 201, r.text
    assert r.json()["data"]["summary"]["rows_valid"] == 5000
    assert peak < 64 * 1024 * 1024
