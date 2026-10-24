"""T-639 / TC-M15-12 — METRIC-DQ-02 (G3, R-DEC-68) + DQ per domain in `dq_domain_scores`."""

import uuid
from datetime import UTC, datetime

import pytest
from sqlalchemy import func, select

from app.core.capture_share import CAPTURE_SOURCES
from app.core.clock import month_key
from app.modules.platform.models import DqDomainScore
from app.modules.platform.services.capture_metrics import capture_share
from app.modules.platform.services.capture_metrics import compute as dq02
from app.modules.platform.services.metrics import compute as dq01
from app.workers.consumers import load_metric_providers
from tests.conftest import ApiClient
from tests.integration.idp.test_lineage_t637 import ROWS, _committed
from tests.integration.idp.test_undo_t638 import _manual_task, _undo

pytestmark = pytest.mark.integration


def _rows(prefix: str) -> list[list[object]]:
    return [[*r[:3], f"{prefix}-{i}", *r[4:]] for i, r in enumerate(ROWS)]


async def _script_share(db, key: str) -> tuple[int, int]:
    """Independent recount (Core SQL on `__table__`) over the registered whitelist tables."""
    captured = total = 0
    for sources in CAPTURE_SOURCES.sources.values():
        for s in sources:
            t = s.model.__table__
            cap = (t.c.data_source.in_(("document", "ai_assisted")) if "data_source" in t.c
                   else t.c.source.in_(("IDP", "IMPORT", "AI")))  # fmt: skip
            month = func.to_char(func.timezone("UTC", t.c.created_at), "YYYY-MM")
            where = [month == key]
            if "deleted_at" in t.c:
                where.append(t.c.deleted_at.is_(None))
            q = select(func.count(), func.count().filter(cap)).select_from(t).where(*where)
            n, c = (await db.execute(q)).one()
            total, captured = total + n, captured + c
    return captured, total


async def test_share_counts_import_not_manual_and_undo_reverts(api: ApiClient, db_session) -> None:
    load_metric_providers()
    assert {"hr", "execution", "finance"} <= set(CAPTURE_SOURCES.sources)
    key = month_key(datetime.now(UTC).date())
    base = await capture_share(db_session, key)
    assert base == await _script_share(db_session, key)
    await _manual_task(api, db_session, uuid.uuid4())  # C1 → MANUAL: total +1, captured +0
    c1 = await capture_share(db_session, key)
    assert (c1[0], c1[1]) == (base[0], base[1] + 1)
    doc_id, refs = await _committed(api, db_session, _rows(uuid.uuid4().hex[:6]))
    after = await capture_share(db_session, key)
    assert after == await _script_share(db_session, key)
    created = len(refs["hr.employees"])
    assert after[0] - c1[0] >= created and after[1] - c1[1] == after[0] - c1[0]
    r = await _undo(api, "HR_MANAGER", str(doc_id), uuid.uuid4())
    assert r.status_code == 200, r.text
    assert await capture_share(db_session, key) == c1  # undo → back to the pre-import share
    reading = (await dq02(db_session, datetime.now(UTC).date()))[0]
    assert reading.code == "METRIC-DQ-02"
    assert reading.value == round(c1[0] / c1[1] * 100, 1)


async def test_no_rows_in_month_is_no_data(db_session) -> None:
    from datetime import date

    (r,) = await dq02(db_session, date(2040, 1, 15))
    assert r.value is None and r.status_reason == "no records created"


async def test_dq01_writes_domain_scores_and_keeps_value(db_session) -> None:
    from datetime import date

    load_metric_providers()
    as_of = date(2026, 10, 20)
    first = await dq01(db_session, as_of)
    rows = (await db_session.execute(select(DqDomainScore).where(
        DqDomainScore.period_key == "2026-10"))).scalars().all()  # fmt: skip
    assert {r.domain for r in rows} >= {"finance", "hr", "execution"}
    mean = round(sum(float(r.score) for r in rows) / len(rows), 1)
    assert first[0].value == mean  # company value = mean of the stored domain scores
    again = await dq01(db_session, as_of)  # upsert, no duplicate rows
    assert again[0].value == first[0].value
    n = (await db_session.execute(select(DqDomainScore).where(
        DqDomainScore.period_key == "2026-10"))).scalars().all()  # fmt: skip
    assert len(n) == len(rows)


async def test_dq01_breakdown_reads_domain_scores(db_session) -> None:
    from datetime import date

    from app.modules.platform.services.metric_query import MetricQueryService

    load_metric_providers()
    await dq01(db_session, date(2026, 10, 20))
    rows = await MetricQueryService(db_session).breakdown("METRIC-DQ-01", "2026-10")
    assert rows and all(t == "domain" and v is not None for t, _, v in rows)
    assert {d for _, d, _ in rows} >= {"finance", "hr"}


def _doc06_rows() -> list[list[object]]:
    from openpyxl import load_workbook

    from tests.unit.test_demo_docs_t641 import DOC06

    ws = load_workbook(DOC06, read_only=True).worksheets[0]
    return [list(r) for r in list(ws.iter_rows(values_only=True))[1:]]


async def test_doc06_import_splits_error_rows(api: ApiClient, db_session, idp_store) -> None:
    """TC-M15-12/15: DOC-06 (scripts/gen_demo_docs.py) → 45 valid + 3 error rows. Codes are
    re-tagged so re-runs on the shared test DB do not collide with an earlier import."""
    from tests.integration.idp.test_import_c5_t635 import _upload, _xlsx

    rows, tag = _doc06_rows(), uuid.uuid4().hex[:5]
    for r in rows:
        r[0] = f"{r[0]}-{tag}" if r[0] else None
        r[3] = f"{r[3]}-{tag}"
        r[5] = f"{r[5]}-{tag}" if r[5] and r[5] != "NV-999" else r[5]
    p = (await _upload(api, "CEO", _xlsx(rows, tag=tag))).json()["data"]
    assert p["suggested_document_type"]["code"] == "HR_ROSTER"
    assert p["summary"] == {"rows_total": 48, "rows_valid": 45, "rows_invalid": 3}
