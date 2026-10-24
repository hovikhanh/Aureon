"""T-641 / TC-M15-15 — demo sample files DOC-06/07/08 in `data/` (08 §4.6, DEC-120).

Reads the committed outputs of `scripts/gen_demo_docs.py` with the runtime parsers (openpyxl,
python-docx, pypdf) — the generator itself needs the `demo` group (reportlab) and asserts its own
totals; byte-identical re-runs are checked by the generator's fixed zip/PDF metadata.
"""

import io
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

DATA = Path(__file__).resolve().parents[4] / "data"
DOC06 = DATA / "DOC-06-danh-sach-nhan-su.xlsx"
DOC07X, DOC07W = DATA / "DOC-07-kpi-okr-q4-2026.xlsx", DATA / "DOC-07-kpi-okr-q4-2026.docx"
DOC08 = DATA / "DOC-08-bao-cao-ket-qua-q3-2026.pdf"


def _rows(path: Path) -> tuple[list[str], list[tuple[object, ...]]]:
    ws = load_workbook(path, read_only=True).worksheets[0]
    head, *rows = list(ws.iter_rows(values_only=True))
    return [str(h) for h in head], rows


def test_doc06_roster_with_intentional_error_rows() -> None:
    head, rows = _rows(DOC06)
    assert head[:9] == ["Mã phòng ban", "Phòng ban", "Chức danh", "Mã NV", "Họ tên",
                        "Quản lý trực tiếp", "Ngày vào làm", "Trạng thái", "Lương"]  # fmt: skip
    codes = {r[3] for r in rows}
    assert len(rows) == 48
    assert sum(int(r[8]) for r in rows[:45]) == 855_000_000  # type: ignore[call-overload]
    bad = rows[45:]
    assert any(r[5] not in codes for r in bad)  # unknown manager
    assert any(str(r[6]).startswith("15/13/") for r in bad)  # wrong date format
    assert any(r[0] is None and r[1] == "Phòng Marketing" for r in bad)  # unknown department
    assert all(r[5] in codes for r in rows[:45] if r[5] is not None)


def test_doc07_xlsx_and_docx_carry_the_same_plan() -> None:
    head, rows = _rows(DOC07X)
    assert head[:3] == ["Kỳ", "Mục tiêu", "Kết quả then chốt"] and "Mã KPI" in head
    assert len(rows) == 11 and {r[0] for r in rows} == {"2026-Q4"}
    assert len({r[1] for r in rows}) == 5
    doc = Document(str(DOC07W))
    krs = [row.cells[2].text for t in doc.tables for row in t.rows[1:]]
    assert sorted(krs) == sorted(str(r[2]) for r in rows)


def test_doc08_period_report_pdf() -> None:
    text = "".join(p.extract_text() or "" for p in PdfReader(str(DOC08)).pages)
    assert "BÁO CÁO KẾT QUẢ QUÝ 3/2026" in text
    assert text.count("KPI-CO-01") == 3 and "76,5 tỷ" in text


def test_ooxml_outputs_have_fixed_timestamps() -> None:
    for path in (DOC06, DOC07X, DOC07W):
        z = zipfile.ZipFile(io.BytesIO(path.read_bytes()))
        assert {i.date_time for i in z.infolist()} == {(2026, 10, 20, 8, 0, 0)}
        assert b"2026-10-20T08:00:00Z" in z.read("docProps/core.xml")
