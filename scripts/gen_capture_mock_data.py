"""Export the demo workbooks DOC-06 / DOC-07 (`data/`, from `scripts/gen_demo_docs.py`) as the
mock BFF C5 content `apps/web/src/mocks/capture-data.json` (T-643): every sheet, cells as text
(blank → ""), so the mock import parses the same rows as the real API (48/45/3 for DOC-06).

Run: uv run python scripts/gen_capture_mock_data.py
"""

import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "apps" / "web" / "src" / "mocks" / "capture-data.json"
FILES = {"hr": "DOC-06-danh-sach-nhan-su.xlsx", "kpi": "DOC-07-kpi-okr-q4-2026.xlsx"}


def cell(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def sheets(name: str) -> list[dict[str, object]]:
    wb = load_workbook(ROOT / "data" / name, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = [[cell(c) for c in r] for r in ws.iter_rows(values_only=True)]
        while rows and not any(rows[-1]):
            rows.pop()
        out.append({"name": ws.title, "rows": rows})
    return out


def main() -> None:
    data = {k: {"file": f, "sheets": sheets(f)} for k, f in FILES.items()}
    lines = ["{"]
    for i, (k, v) in enumerate(data.items()):
        lines.append(f' "{k}": {{"file": {json.dumps(v["file"])}, "sheets": [')
        for j, s in enumerate(v["sheets"]):  # type: ignore[arg-type]
            lines.append(f'  {{"name": {json.dumps(s["name"], ensure_ascii=False)}, "rows": [')
            rows = [json.dumps(r, ensure_ascii=False) for r in s["rows"]]
            lines.append(",\n".join(f"   {r}" for r in rows))
            lines.append("  ]}" + ("," if j < len(v["sheets"]) - 1 else ""))  # type: ignore[arg-type]
        lines.append(" ]}" + ("," if i < len(data) - 1 else ""))
    lines.append("}")
    OUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
